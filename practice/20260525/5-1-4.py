"""
折れ線グラフの作成には、LineChartオブジェクトを作成して、add_dataメソッドの引数にReferenceオブジェクトを設定
WorkSheetオブジェクトのadd_chartメソッドで、lineChartオブジェクトとグラフを追加するセル番地を指定して、
シートに折れ線グラフを作成する
"""
from openpyxl import load_workbook
from openpyxl.chart import LineChart,Reference

wb = load_workbook("practice/20260525/売上まとめ.xlsx")
ws = wb.active
data = Reference(ws,min_col=1, min_row=2, max_col=4, max_row=4)
label = Reference(ws,min_col=2, min_row=1,max_col=4,max_row=1)
Line = LineChart()
# from_rowsのキーワード引数trueにするとデータを1ずつ読み込む。falseにすると列ごとに読み込む
# 各行の先頭（1列目）をラベルに使う
Line.add_data(data,from_rows=True,titles_from_data=True)
Line.set_categories(label)
Line.x_axis.title = "月"
Line.y_axis.title ="売上金額"
Line.title ="売上別売上金額"
ws.add_chart(Line,"F1")
wb.save("practice/20260525/売上まとめ(折れ線グラフ).xlsx")


