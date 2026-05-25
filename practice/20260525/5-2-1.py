from openpyxl import load_workbook
from openpyxl.chart import BarChart,Reference

wb = load_workbook("practice/20260525/売上.xlsx")
ws = wb.active
data = Reference(ws,min_col=2, min_row=1,max_col=2, max_row=4)
label = Reference(ws,min_col=1,min_row=2, max_col=1, max_row=4)
bar = BarChart()
bar.type ="col"
# 各列の先頭（1行目）をタイトルとして使う
bar.add_data(data,titles_from_data=True)
bar.set_categories(label)
bar.x_axis.title = "品名"
bar.y_axis.title ="売上金額"
bar.title = "売上別金額"
bar.width = 12
bar.height =6
ws.add_chart(bar,"D1")
wb.save("practice/20260525/売上(棒グラフサイズ変更後).xlsx")