from openpyxl import load_workbook
from openpyxl.chart import BarChart,Reference

wb = load_workbook("practice/20260525/売上.xlsx")
sheet_list = wb.worksheets
for ws in sheet_list:
    data = Reference(ws,min_col=2, min_row=1, max_col= 2, max_row=4)
    label = Reference(ws,min_col=1,min_row=2, max_col=1, max_row=4)
    bar = BarChart()
    bar.type = "col"
    bar.add_data(data,titles_from_data=True)
    bar.set_categories(label)
    bar.title = ws.title
    bar.width = 7
    bar.height =7
    bar.legend.position = "b"
    bar.varyColors = True
    bar.y_axis.scaling.max = 15000
    ws.add_chart(bar,"A6")
wb.save("practice/20260525/売上(棒グラフ複数作成後).xlsx")